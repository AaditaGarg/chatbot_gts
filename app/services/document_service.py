"""Document upload, extraction, chunking, and indexing."""

import csv
import io
import logging
import os
import uuid
from pathlib import Path
from typing import Optional

from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.orm import Session

from config.settings import ALLOWED_UPLOAD_EXTENSIONS, CHUNK_OVERLAP, CHUNK_SIZE, MAX_UPLOAD_SIZE_MB, UPLOAD_DIR
from database.postgres import Document
from services.rag_service import RAGService

logger = logging.getLogger(__name__)


class DocumentService:
    """Handle file uploads and index them in ChromaDB."""

    def __init__(self, rag_service: RAGService):
        self.rag = rag_service
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
        )
        os.makedirs(UPLOAD_DIR, exist_ok=True)

    def validate_file(self, filename: str, file_size: int) -> None:
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_UPLOAD_EXTENSIONS:
            raise ValueError(f"Unsupported file type '{ext}'. Allowed: {', '.join(ALLOWED_UPLOAD_EXTENSIONS)}")
        max_bytes = MAX_UPLOAD_SIZE_MB * 1024 * 1024
        if file_size > max_bytes:
            raise ValueError(f"File too large. Maximum size is {MAX_UPLOAD_SIZE_MB} MB.")

    def extract_text(self, filename: str, content: bytes) -> str:
        ext = Path(filename).suffix.lower()
        if ext == ".txt":
            return content.decode("utf-8", errors="replace")
        if ext == ".csv":
            return self._extract_csv(content)
        if ext == ".pdf":
            return self._extract_pdf(content)
        if ext == ".docx":
            return self._extract_docx(content)
        raise ValueError(f"Unsupported file type: {ext}")

    def _extract_csv(self, content: bytes) -> str:
        text = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [" | ".join(row) for row in reader]
        return "\n".join(rows)

    def _extract_pdf(self, content: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support requires 'pypdf'. Install with: pip install pypdf") from exc

        reader = PdfReader(io.BytesIO(content))
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[Page {i}]\n{page_text}")
        return "\n\n".join(pages)

    def _extract_docx(self, content: bytes) -> str:
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("DOCX support requires 'python-docx'. Install with: pip install python-docx") from exc

        document = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())

    def ingest_file(
        self,
        db: Session,
        user_id: uuid.UUID,
        session_id: Optional[uuid.UUID],
        filename: str,
        content: bytes,
    ) -> Document:
        self.validate_file(filename, len(content))
        text = self.extract_text(filename, content).strip()
        if not text:
            raise ValueError("Could not extract any text from the uploaded file.")

        doc_record = Document(
            user_id=user_id,
            conversation_id=session_id,
            filename=filename,
            file_type=Path(filename).suffix.lower(),
            file_size=len(content),
        )
        db.add(doc_record)
        db.commit()
        db.refresh(doc_record)

        save_path = os.path.join(UPLOAD_DIR, f"{doc_record.document_id}{doc_record.file_type}")
        with open(save_path, "wb") as f:
            f.write(content)

        splits = self.splitter.split_text(text)
        ids = [f"{doc_record.document_id}_{i}" for i in range(len(splits))]
        metadatas = [
            {
                "user_id": str(user_id),
                "document_id": str(doc_record.document_id),
                "session_id": str(session_id) if session_id else "",
                "filename": filename,
                "chunk_index": i,
            }
            for i in range(len(splits))
        ]

        chunk_count = self.rag.add_documents(chunks=splits, ids=ids, metadatas=metadatas)
        doc_record.chunk_count = chunk_count
        db.commit()
        db.refresh(doc_record)

        logger.info("Indexed %d chunks for document %s", chunk_count, doc_record.document_id)
        return doc_record
